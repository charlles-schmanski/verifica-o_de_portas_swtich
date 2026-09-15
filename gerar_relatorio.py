"""
Gera um relatorio tecnico em Word a partir de arquivos JSON historicos
de portas de switches.

Exemplos:
    python gerar_relatorio_portas.py
    python gerar_relatorio_portas.py -i historico -o Relatorio_Portas.docx
    python gerar_relatorio_portas.py -i historico -p "portas_*.json"

Dependencia:
    pip install python-docx
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


@dataclass
class Observacao:
    data_hora: datetime
    switch: str
    host: str
    device_type: str
    porta: str
    status: str
    raw_status: str
    arquivo: str


def argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Gera relatorio tecnico de utilizacao de portas a partir de historicos JSON."
    )
    parser.add_argument(
        "-i", "--entrada", default="historico",
        help="Pasta que contem os JSONs. Subpastas tambem sao pesquisadas. Default: historico",
    )
    parser.add_argument(
        "-p", "--padrao", default="portas_*.json",
        help="Padrao dos arquivos JSON. Default: portas_*.json",
    )
    parser.add_argument(
        "-o", "--saida", default="Relatorio_Tecnico_Portas_Switches.docx",
        help="Arquivo DOCX de saida.",
    )
    parser.add_argument(
        "--titulo", default="RELATORIO TECNICO",
        help="Titulo principal do documento.",
    )
    parser.add_argument(
        "--data",
        help="Seleciona uma data especifica no formato AAAA-MM-DD. Exemplo: 2026-09-04",
    )
    parser.add_argument(
        "--inicio",
        help="Data inicial no formato AAAA-MM-DD. Deve ser usada junto com --fim.",
    )
    parser.add_argument(
        "--fim",
        help="Data final no formato AAAA-MM-DD. Deve ser usada junto com --inicio.",
    )
    parser.add_argument(
        "--listar-datas",
        action="store_true",
        help="Lista as datas disponiveis nos arquivos e encerra.",
    )
    return parser.parse_args()


def carregar_arquivos(pasta: Path, padrao: str) -> tuple[list[Observacao], list[dict[str, Any]]]:
    arquivos = sorted(pasta.rglob(padrao))
    if not arquivos:
        raise FileNotFoundError(
            f"Nenhum arquivo encontrado em '{pasta}' usando o padrao '{padrao}'."
        )

    observacoes: list[Observacao] = []
    resumos_arquivos: list[dict[str, Any]] = []

    for arquivo in arquivos:
        with arquivo.open("r", encoding="utf-8") as f:
            dados = json.load(f)

        if not isinstance(dados, list):
            print(f"[AVISO] Ignorado, JSON nao e uma lista: {arquivo}")
            continue

        registros_validos = 0
        datas: list[datetime] = []
        switches: set[str] = set()

        for item in dados:
            try:
                dt = datetime.fromisoformat(str(item["data_hora"]))
                obs = Observacao(
                    data_hora=dt,
                    switch=str(item["switch"]),
                    host=str(item.get("host", "")),
                    device_type=str(item.get("device_type", "")),
                    porta=str(item["porta"]),
                    status=str(item.get("status", "UNKNOWN")).upper(),
                    raw_status=str(item.get("raw_status", "")),
                    arquivo=arquivo.name,
                )
            except (KeyError, TypeError, ValueError) as erro:
                print(f"[AVISO] Registro invalido em {arquivo}: {erro}")
                continue

            observacoes.append(obs)
            datas.append(dt)
            switches.add(obs.switch)
            registros_validos += 1

        if registros_validos:
            resumos_arquivos.append(
                {
                    "arquivo": arquivo.name,
                    "inicio": min(datas),
                    "fim": max(datas),
                    "switches": sorted(switches),
                    "registros": registros_validos,
                }
            )

    if not observacoes:
        raise ValueError("Nenhuma observacao valida foi encontrada nos JSONs.")

    observacoes.sort(key=lambda x: (x.data_hora, x.switch, chave_natural(x.porta)))
    resumos_arquivos.sort(key=lambda x: x["inicio"])
    return observacoes, resumos_arquivos


def chave_natural(texto: str) -> list[Any]:
    return [int(x) if x.isdigit() else x.lower() for x in re.split(r"(\d+)", texto)]


def formatar_duracao(inicio: datetime, fim: datetime) -> str:
    segundos = max(0, int((fim - inicio).total_seconds()))
    dias, resto = divmod(segundos, 86400)
    horas, resto = divmod(resto, 3600)
    minutos, _ = divmod(resto, 60)
    return f"{dias}d {horas:02d}h {minutos:02d}min"


def formatar_data(dt: datetime) -> str:
    return dt.strftime("%d/%m/%Y %H:%M:%S")


def configurar_margens_celula(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margem, valor in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elemento = tc_mar.find(qn(f"w:{margem}"))
        if elemento is None:
            elemento = OxmlElement(f"w:{margem}")
            tc_mar.append(elemento)
        elemento.set(qn("w:w"), str(valor))
        elemento.set(qn("w:type"), "dxa")


def aplicar_cor_celula(cell, cor: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), cor)
    tc_pr.append(shd)


def estilizar_tabela(tabela, larguras: list[float] | None = None) -> None:
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.style = "Table Grid"

    cabecalho = tabela.rows[0]
    tr_pr = cabecalho._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    tr_pr.append(repetir)

    for cell in cabecalho.cells:
        aplicar_cor_celula(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        configurar_margens_celula(cell, 110, 100, 110, 100)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)

    for row in tabela.rows[1:]:
        for indice, cell in enumerate(row.cells):
            configurar_margens_celula(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                if indice > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8)

    if larguras:
        for row in tabela.rows:
            for indice, largura in enumerate(larguras):
                row.cells[indice].width = Inches(largura)


def adicionar_titulo_secao(doc: Document, texto: str, nivel: int = 1):
    p = doc.add_heading(texto, level=nivel)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    return p


def adicionar_nota(doc: Document, titulo: str, texto: str) -> None:
    tabela = doc.add_table(rows=1, cols=1)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tabela.cell(0, 0)
    aplicar_cor_celula(cell, "EAF2F8")
    configurar_margens_celula(cell, 140, 160, 140, 160)
    p = cell.paragraphs[0]
    run = p.add_run(titulo + " ")
    run.bold = True
    run.font.color.rgb = RGBColor(31, 78, 120)
    p.add_run(texto)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def agrupar_portas(observacoes: list[Observacao]):
    grupos: defaultdict[tuple[str, str, str], list[Observacao]] = defaultdict(list)
    for obs in observacoes:
        grupos[(obs.switch, obs.host, obs.porta)].append(obs)
    for chave in grupos:
        grupos[chave].sort(key=lambda x: x.data_hora)
    return grupos


def filtrar_observacoes_por_data(
    observacoes: list[Observacao],
    resumos_arquivos: list[dict[str, Any]],
    data: str | None = None,
    inicio: str | None = None,
    fim: str | None = None,
) -> tuple[list[Observacao], list[dict[str, Any]]]:
    if data and (inicio or fim):
        raise ValueError("Use --data ou o intervalo --inicio/--fim, nao ambos.")

    if bool(inicio) != bool(fim):
        raise ValueError("Os parametros --inicio e --fim devem ser usados juntos.")

    if not data and not inicio and not fim:
        return observacoes, resumos_arquivos

    try:
        if data:
            data_inicial = datetime.strptime(data, "%Y-%m-%d").date()
            data_final = data_inicial
        else:
            data_inicial = datetime.strptime(inicio, "%Y-%m-%d").date()
            data_final = datetime.strptime(fim, "%Y-%m-%d").date()
    except ValueError as erro:
        raise ValueError("Data invalida. Use o formato AAAA-MM-DD, por exemplo 2026-09-04.") from erro

    if data_inicial > data_final:
        raise ValueError("A data inicial nao pode ser posterior a data final.")

    filtradas = [
        obs for obs in observacoes
        if data_inicial <= obs.data_hora.date() <= data_final
    ]

    if not filtradas:
        raise ValueError(
            f"Nenhuma coleta encontrada entre {data_inicial.isoformat()} e {data_final.isoformat()}."
        )

    arquivos_usados = {obs.arquivo for obs in filtradas}
    resumos_filtrados = [
        resumo for resumo in resumos_arquivos
        if resumo["arquivo"] in arquivos_usados
    ]

    return filtradas, resumos_filtrados


def listar_datas_disponiveis(observacoes: list[Observacao]) -> None:
    datas = sorted({obs.data_hora.date() for obs in observacoes})
    print("Datas disponiveis:")
    for data in datas:
        quantidade = sum(1 for obs in observacoes if obs.data_hora.date() == data)
        switches = sorted({
            obs.switch for obs in observacoes if obs.data_hora.date() == data
        })
        print(
            f"  {data.isoformat()} | {quantidade} registros | "
            + ", ".join(switches)
        )


def criar_relatorio(
    observacoes: list[Observacao],
    resumos_arquivos: list[dict[str, Any]],
    saida: Path,
    titulo: str,
) -> None:
    doc = Document()
    secao = doc.sections[0]
    secao.top_margin = Inches(0.65)
    secao.bottom_margin = Inches(0.65)
    secao.left_margin = Inches(0.70)
    secao.right_margin = Inches(0.70)

    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)
    for nome in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[nome].font.name = "Aptos Display"
        doc.styles[nome].font.color.rgb = RGBColor(31, 78, 120)

    cabecalho = secao.header.paragraphs[0]
    cabecalho.text = "RELATORIO TECNICO | MONITORAMENTO DE PORTAS"
    cabecalho.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in cabecalho.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = rodape.add_run("Documento gerado automaticamente a partir do historico JSON")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

    inicio_periodo = min(x.data_hora for x in observacoes)
    fim_periodo = max(x.data_hora for x in observacoes)
    switches = sorted({x.switch for x in observacoes})
    grupos = agrupar_portas(observacoes)

    # Capa
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(27)
    r.font.color.rgb = RGBColor(31, 78, 120)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Analise de utilizacao e variacao de portas de switches")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(70, 90, 110)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.add_run("Periodo observado\n").bold = True
    p.add_run(f"{formatar_data(inicio_periodo)} a {formatar_data(fim_periodo)}")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.add_run("Equipamentos analisados\n").bold = True
    p.add_run(" | ".join(switches))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run(
        "Finalidade: identificar portas ativas, alteracoes de estado e janela observada de utilizacao, separada por switch."
    )
    r.italic = True
    r.font.size = Pt(10)
    doc.add_page_break()

    adicionar_titulo_secao(doc, "1. Objetivo e escopo")
    doc.add_paragraph(
        "Este relatorio compara os arquivos historicos de coleta de portas fisicas. "
        "A analise apresenta o estado UP/DOWN por equipamento, as diferencas entre "
        "as coletas e a janela entre a primeira e a ultima observacao UP de cada porta."
    )
    adicionar_nota(
        doc,
        "Limitacao metodologica:",
        "as coletas sao fotografias pontuais. Uma porta UP em duas coletas foi observada "
        "ativa nesses instantes, mas isso nao comprova atividade continua entre eles. "
        "A coluna de tempo deve ser interpretada como janela observada, e nao uptime exato.",
    )

    adicionar_titulo_secao(doc, "2. Arquivos analisados")
    tabela = doc.add_table(rows=1, cols=4)
    for i, cab in enumerate(("Arquivo", "Data/hora", "Switches presentes", "Registros")):
        tabela.cell(0, i).text = cab
    for resumo in resumos_arquivos:
        row = tabela.add_row().cells
        row[0].text = resumo["arquivo"]
        row[1].text = f'{formatar_data(resumo["inicio"])} a {formatar_data(resumo["fim"])}'
        row[2].text = ", ".join(resumo["switches"])
        row[3].text = str(resumo["registros"])
    estilizar_tabela(tabela, [1.55, 2.15, 2.45, 0.75])

    adicionar_titulo_secao(doc, "3. Resumo por coleta")
    tabela = doc.add_table(rows=1, cols=6)
    for i, cab in enumerate(("Coleta", "Switches", "Portas", "UP", "DOWN", "% UP")):
        tabela.cell(0, i).text = cab
    for resumo in resumos_arquivos:
        registros = [x for x in observacoes if x.arquivo == resumo["arquivo"]]
        contagem = Counter(x.status for x in registros)
        total = len(registros)
        row = tabela.add_row().cells
        valores = (
            formatar_data(resumo["inicio"]),
            len(set(x.switch for x in registros)),
            total,
            contagem["UP"],
            contagem["DOWN"],
            f'{(contagem["UP"] / total * 100):.1f}%' if total else "0.0%",
        )
        for i, valor in enumerate(valores):
            row[i].text = str(valor)
    estilizar_tabela(tabela, [1.55, 0.8, 0.75, 0.6, 0.7, 0.8])

    adicionar_titulo_secao(doc, "4. Analise por switch")
    mudancas: list[tuple[str, str, Observacao, Observacao]] = []

    for switch in switches:
        adicionar_titulo_secao(doc, switch, 2)
        obs_switch = [x for x in observacoes if x.switch == switch]
        host = obs_switch[0].host
        tipo = obs_switch[0].device_type
        doc.add_paragraph(f"Host: {host} | Tipo: {tipo}")

        tabela = doc.add_table(rows=1, cols=5)
        for i, cab in enumerate(("Coleta", "Total", "UP", "DOWN", "% UP")):
            tabela.cell(0, i).text = cab

        arquivos_switch = sorted(
            {x.arquivo for x in obs_switch},
            key=lambda arq: min(x.data_hora for x in obs_switch if x.arquivo == arq),
        )
        for arquivo in arquivos_switch:
            registros = [x for x in obs_switch if x.arquivo == arquivo]
            contagem = Counter(x.status for x in registros)
            total = len(registros)
            row = tabela.add_row().cells
            valores = (
                formatar_data(min(x.data_hora for x in registros)),
                total,
                contagem["UP"],
                contagem["DOWN"],
                f'{(contagem["UP"] / total * 100):.1f}%' if total else "0.0%",
            )
            for i, valor in enumerate(valores):
                row[i].text = str(valor)
        estilizar_tabela(tabela, [1.7, 0.7, 0.6, 0.7, 0.8])

        portas_ativas = []
        for (sw, _host, porta), lista in grupos.items():
            if sw != switch:
                continue
            ups = [x for x in lista if x.status == "UP"]
            if not ups:
                continue
            duracao = (
                formatar_duracao(ups[0].data_hora, ups[-1].data_hora)
                if len(ups) > 1 else "Somente 1 observacao"
            )
            sequencia = " -> ".join(x.status for x in lista)
            portas_ativas.append(
                (porta, len(lista), len(ups), duracao, sequencia, lista[-1].status)
            )

            for anterior, posterior in zip(lista, lista[1:]):
                if anterior.status != posterior.status:
                    mudancas.append((switch, porta, anterior, posterior))

        portas_ativas.sort(key=lambda x: chave_natural(x[0]))
        doc.add_paragraph("Portas observadas como ativas no periodo:")
        tabela = doc.add_table(rows=1, cols=6)
        cabecalhos = (
            "Porta", "Observacoes", "UP", "Janela entre 1º e ultimo UP",
            "Sequencia", "Estado final",
        )
        for i, cab in enumerate(cabecalhos):
            tabela.cell(0, i).text = cab
        for valores in portas_ativas:
            row = tabela.add_row().cells
            for i, valor in enumerate(valores):
                row[i].text = str(valor)
            aplicar_cor_celula(row[5], "E2F0D9" if valores[-1] == "UP" else "FCE4D6")
        estilizar_tabela(tabela, [1.6, 0.75, 0.45, 1.6, 1.35, 0.75])

        persistentes = sum(1 for x in portas_ativas if x[4].replace(" ", "").split("->").count("UP") == x[1])
        doc.add_paragraph(
            f"Foram identificadas {len(portas_ativas)} portas com pelo menos uma observacao UP. "
            f"Dessas, {persistentes} permaneceram UP em todas as observacoes disponiveis do equipamento."
        )

    adicionar_titulo_secao(doc, "5. Diferencas entre as coletas")
    tabela = doc.add_table(rows=1, cols=6)
    for i, cab in enumerate(
        ("Switch", "Porta", "Estado anterior", "Data anterior", "Estado posterior", "Data posterior")
    ):
        tabela.cell(0, i).text = cab

    if mudancas:
        for switch, porta, anterior, posterior in sorted(
            mudancas, key=lambda x: (x[0], chave_natural(x[1]), x[2].data_hora)
        ):
            row = tabela.add_row().cells
            valores = (
                switch, porta, anterior.status, formatar_data(anterior.data_hora),
                posterior.status, formatar_data(posterior.data_hora),
            )
            for i, valor in enumerate(valores):
                row[i].text = str(valor)
    else:
        row = tabela.add_row().cells
        row[0].merge(row[-1]).text = "Nenhuma mudanca de estado foi identificada."
    estilizar_tabela(tabela, [1.25, 1.8, 0.8, 1.25, 0.8, 1.25])

    adicionar_titulo_secao(doc, "6. Conclusao")
    total_portas = len({(x.switch, x.host, x.porta) for x in observacoes})
    total_mudancas = len(mudancas)
    doc.add_paragraph(
        f"Foram analisados {len(resumos_arquivos)} arquivos, {len(switches)} switches e "
        f"{total_portas} portas unicas. Foram identificadas {total_mudancas} mudancas de estado "
        "entre observacoes consecutivas. O historico permite comparar disponibilidade e estimar "
        "janelas observadas de atividade, respeitando a frequencia das coletas como limite de precisao."
    )

    adicionar_titulo_secao(doc, "Anexo A. Portas UP na coleta mais recente")
    ultima_data_por_switch = {
        sw: max(x.data_hora for x in observacoes if x.switch == sw) for sw in switches
    }
    for switch in switches:
        ultima_data = ultima_data_por_switch[switch]
        portas = sorted(
            {
                x.porta for x in observacoes
                if x.switch == switch and x.data_hora == ultima_data and x.status == "UP"
            },
            key=chave_natural,
        )
        doc.add_paragraph(f"{switch}: {len(portas)} porta(s) UP", style="Heading 3")
        doc.add_paragraph(", ".join(portas) if portas else "Nenhuma porta UP.")

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(saida)


def main() -> int:
    args = argumentos()
    pasta = Path(args.entrada)
    saida = Path(args.saida)

    try:
        observacoes, resumos = carregar_arquivos(pasta, args.padrao)

        if args.listar_datas:
            listar_datas_disponiveis(observacoes)
            return 0

        observacoes, resumos = filtrar_observacoes_por_data(
            observacoes,
            resumos,
            data=args.data,
            inicio=args.inicio,
            fim=args.fim,
        )

        criar_relatorio(observacoes, resumos, saida, args.titulo)
    except Exception as erro:
        print(f"[ERRO] {erro}")
        return 1

    print(f"[OK] Relatorio gerado: {saida.resolve()}")
    print(f"[OK] Arquivos analisados: {len(resumos)}")
    print(f"[OK] Registros analisados: {len(observacoes)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
